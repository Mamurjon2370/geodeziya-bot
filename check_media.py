import docx
import os
import sys
import zipfile
import xml.etree.ElementTree as ET

sys.stdout.reconfigure(encoding='utf-8')

def check_images(docx_path):
    with zipfile.ZipFile(docx_path, 'r') as z:
        media_files = [f for f in z.namelist() if f.startswith('word/media/')]
        print(f"{docx_path} contains {len(media_files)} media files/images.")

def parse_docx(filename):
    doc = docx.Document(filename)
    print(f"\n==============================")
    print(f"Parsing: {filename}")
    print(f"Total paragraphs: {len(doc.paragraphs)}")
    
    questions = []
    current_q = None
    
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if not text:
            continue
            
        # Check if text is "Тест" or header
        if text.lower() == "тест":
            continue
            
        # Check if this paragraph has yellow highlight or other highlight
        is_highlighted = False
        for r in p.runs:
            if r.font.highlight_color: # e.g. WD_COLOR_INDEX.YELLOW
                is_highlighted = True
                break
            # also check xml if highlight exists in r._element
            if 'w:highlight' in r._element.xml:
                is_highlighted = True
                break
        
        # Check if it ends with '?' or starts with Roman numeral like 'I. ' or looks like a question
        # Or let's see how questions are distinguished
        # In the inspect output:
        # P8: Ерни масофадан туриб зондлаш кандай жараён ? -> No highlight
        # P10: (yellow) Ер сатхи хакидаги...
        # P11: Ер сатхи хакидаги...
        # P12: планета сатхи...
        
    print("Done inspection preview.")

if __name__ == "__main__":
    check_images("савол жавоблар 1.docx")
    check_images("савол жавоблар 2.docx")

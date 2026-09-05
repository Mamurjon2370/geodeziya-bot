import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

def inspect_tables(filename):
    doc = docx.Document(filename)
    print(f"=== Tables in {filename} (Total: {len(doc.tables)}) ===")
    for i, t in enumerate(doc.tables[:15]):
        print(f"\n--- Table #{i+1} (rows: {len(t.rows)}, cols: {len(t.columns)}) ---")
        for r_idx, row in enumerate(t.rows):
            cell_vals = []
            for c_idx, cell in enumerate(row.cells):
                # check if cell has images
                imgs = cell._element.xpath('.//a:blip/@r:embed')
                cell_vals.append(f"[{c_idx}]: '{cell.text.strip()}' (imgs={imgs})")
            print(f"  Row {r_idx}: {' | '.join(cell_vals)}")

if __name__ == "__main__":
    inspect_tables("савол жавоблар 1.docx")

import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

def inspect_table_styles(filename):
    doc = docx.Document(filename)
    print(f"=== Checking first 10 tables in {filename} ===")
    for i, t in enumerate(doc.tables[:10]):
        print(f"\n--- Table #{i+1} ---")
        for r_idx, row in enumerate(t.rows):
            for c_idx, cell in enumerate(row.cells):
                for p_idx, p in enumerate(cell.paragraphs):
                    text = p.text.strip()
                    if text:
                        hl = []
                        for r in p.runs:
                            if r.font.highlight_color or 'w:highlight' in r._element.xml:
                                hl.append((r.text, str(r.font.highlight_color)))
                        print(f"  Cell[{c_idx}] P{p_idx}: text='{text}' | hl={hl}")

if __name__ == "__main__":
    inspect_table_styles("савол жавоблар 1.docx")

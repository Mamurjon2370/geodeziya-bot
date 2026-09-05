import docx
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

def inspect_doc(filename, output_file):
    with open(output_file, "w", encoding="utf-8") as out:
        out.write(f"=== Inspecting {filename} ===\n")
        doc = docx.Document(filename)
        out.write(f"Total paragraphs: {len(doc.paragraphs)}\n")
        out.write(f"Total tables: {len(doc.tables)}\n\n")
        
        out.write("--- First 50 paragraphs ---\n")
        for i, p in enumerate(doc.paragraphs[:50]):
            text = p.text.strip()
            if text:
                run_info = []
                for r in p.runs:
                    color = None
                    try:
                        if r.font and r.font.color:
                            color = r.font.color.rgb
                    except:
                        pass
                    if r.bold or r.font.underline or color or r.font.highlight_color:
                        run_info.append(f"'{r.text}'(bold={r.bold}, under={r.font.underline}, color={color}, highlight={r.font.highlight_color})")
                out.write(f"P{i}: {text}\n  -> Styled runs: {run_info}\n")

        if doc.tables:
            out.write(f"\n--- First Table sample (rows: {len(doc.tables[0].rows)}) ---\n")
            for r_idx, row in enumerate(doc.tables[0].rows[:10]):
                cells = [c.text.strip().replace('\n', ' ') for c in row.cells]
                out.write(f"Row {r_idx}: {cells}\n")

if __name__ == "__main__":
    if os.path.exists("савол жавоблар 1.docx"):
        inspect_doc("савол жавоблар 1.docx", "inspect_output_1.txt")
    if os.path.exists("савол жавоблар 2.docx"):
        inspect_doc("савол жавобlar 2.docx" if not os.path.exists("савол жавоблар 2.docx") else "савол жавоблар 2.docx", "inspect_output_2.txt")
    print("Inspection finished.")
